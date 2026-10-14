"""Testes de ponta a ponta do motor de análise."""

from __future__ import annotations

import os
import threading
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.engine import AnalysisEngine, EngineCallbacks
from core.models import AnalysisConfig, ProgressSnapshot
from utils.formatter import build_report_filename


def _config(source: Path, destino: Path, **kwargs: object) -> AnalysisConfig:
    return AnalysisConfig(
        source=str(source),
        algorithms=list(kwargs.pop("algorithms", ["SHA-256"])),  # type: ignore[arg-type]
        output_path=destino / build_report_filename(),
        **kwargs,  # type: ignore[arg-type]
    )


def test_analise_completa(sample_tree: Path, tmp_path: Path) -> None:
    """A análise processa todos os arquivos e gera o relatório."""
    destino = tmp_path / "relatorios"
    resumo = AnalysisEngine(_config(sample_tree, destino)).run()

    assert resumo.files_found == 5
    assert resumo.files_processed == 5
    assert resumo.files_ok == 5
    assert resumo.total_errors == 0
    assert not resumo.cancelled
    assert resumo.report_path is not None
    assert resumo.report_path.exists()
    assert resumo.bytes_processed == resumo.bytes_found


def test_quantidade_de_linhas_corresponde_aos_arquivos(sample_tree: Path, tmp_path: Path) -> None:
    """A planilha deve conter exatamente uma linha por arquivo encontrado."""
    resumo = AnalysisEngine(_config(sample_tree, tmp_path)).run()

    workbook = load_workbook(resumo.report_path, read_only=True)
    try:
        planilha = workbook["Arquivos"]
        linhas = [linha for linha in planilha.iter_rows(min_row=2, values_only=True)]
    finally:
        workbook.close()

    assert len(linhas) == resumo.files_found == 5
    # Coluna 9 (índice 8) contém o hash; coluna 10 (índice 9) o status.
    assert all(linha[9] == "OK" for linha in linhas)
    assert all(len(str(linha[8])) == 64 for linha in linhas)


def test_hash_da_planilha_confere_com_o_arquivo(sample_tree: Path, tmp_path: Path) -> None:
    """O hash gravado precisa corresponder ao conteúdo real do arquivo."""
    import hashlib

    resumo = AnalysisEngine(_config(sample_tree, tmp_path)).run()

    workbook = load_workbook(resumo.report_path, read_only=True)
    try:
        registros = {
            linha[3]: linha[8]
            for linha in workbook["Arquivos"].iter_rows(min_row=2, values_only=True)
        }
    finally:
        workbook.close()

    for caminho, hash_gravado in registros.items():
        esperado = hashlib.sha256(Path(caminho).read_bytes()).hexdigest().upper()
        assert hash_gravado == esperado, caminho


def test_multiplos_algoritmos_simultaneos(sample_tree: Path, tmp_path: Path) -> None:
    """Com vários algoritmos, cada um recebe a sua própria coluna."""
    config = _config(sample_tree, tmp_path, algorithms=["MD5", "SHA-256"])
    resumo = AnalysisEngine(config).run()

    workbook = load_workbook(resumo.report_path, read_only=True)
    try:
        planilha = workbook["Arquivos"]
        cabecalho = [celula.value for celula in next(planilha.iter_rows(max_row=1))]
        primeira = next(planilha.iter_rows(min_row=2, max_row=2, values_only=True))
    finally:
        workbook.close()

    assert "Hash MD5" in cabecalho and "Hash SHA-256" in cabecalho
    assert len(str(primeira[cabecalho.index("Hash MD5")])) == 32
    assert len(str(primeira[cabecalho.index("Hash SHA-256")])) == 64


def test_deteccao_de_duplicados(sample_tree: Path, tmp_path: Path) -> None:
    """Arquivos com o mesmo conteúdo formam um grupo de duplicados."""
    resumo = AnalysisEngine(_config(sample_tree, tmp_path)).run()

    assert resumo.duplicate_groups == 1
    assert resumo.duplicate_files == 3

    workbook = load_workbook(resumo.report_path, read_only=True)
    try:
        linhas = list(
            workbook["Hashes Duplicados"].iter_rows(min_row=2, values_only=True)
        )
    finally:
        workbook.close()

    assert len(linhas) == 3
    assert {linha[0] for linha in linhas} == {1}
    assert len({linha[1] for linha in linhas}) == 1


def test_duplicados_desativados(sample_tree: Path, tmp_path: Path) -> None:
    """A detecção de duplicados pode ser desligada pelo usuário."""
    config = _config(sample_tree, tmp_path, detect_duplicates=False)
    resumo = AnalysisEngine(config).run()

    assert resumo.duplicate_groups == 0
    workbook = load_workbook(resumo.report_path, read_only=True)
    try:
        primeira = next(workbook["Hashes Duplicados"].iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()
    assert "desativada" in str(primeira[1]).lower()


def test_erro_de_leitura_nao_encerra_a_analise(
    sample_tree: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Um arquivo ilegível é registrado como erro e a análise prossegue."""
    open_original = open

    def open_seletivo(arquivo, *args, **kwargs):  # type: ignore[no-untyped-def]
        if "foto 01" in str(arquivo):
            raise PermissionError(13, "Permission denied")
        return open_original(arquivo, *args, **kwargs)

    monkeypatch.setattr("core.hasher.open", open_seletivo, raising=False)
    resumo = AnalysisEngine(_config(sample_tree, tmp_path)).run()

    assert resumo.files_found == 5
    assert resumo.files_processed == 5
    assert resumo.files_error == 1
    assert resumo.files_ok == 4
    assert resumo.report_path is not None and resumo.report_path.exists()

    workbook = load_workbook(resumo.report_path, read_only=True)
    try:
        erros = list(workbook["Erros"].iter_rows(min_row=2, values_only=True))
        status = [
            linha[9] for linha in workbook["Arquivos"].iter_rows(min_row=2, values_only=True)
        ]
    finally:
        workbook.close()

    assert len(erros) == 1
    assert erros[0][3] == "Acesso negado"
    assert any(str(item).startswith("ERRO - Acesso negado") for item in status)


def test_cancelamento_salva_resultados_parciais(sample_tree: Path, tmp_path: Path) -> None:
    """O cancelamento gera relatório com os resultados já obtidos."""
    # batch_size=1 garante um ponto de cancelamento após cada arquivo.
    engine = AnalysisEngine(_config(sample_tree, tmp_path, batch_size=1))
    processados: list[int] = []

    def ao_progredir(snapshot: ProgressSnapshot) -> None:
        # Cancela assim que o primeiro arquivo terminar de ser processado.
        if snapshot.files_processed >= 1 and not engine.cancelled:
            processados.append(snapshot.files_processed)
            engine.cancel()

    engine.callbacks = EngineCallbacks(on_progress=ao_progredir)
    resumo = engine.run()

    assert resumo.cancelled
    assert resumo.files_found == 5
    assert resumo.files_processed < 5
    assert resumo.files_not_processed > 0
    assert resumo.report_path is not None and resumo.report_path.exists()

    workbook = load_workbook(resumo.report_path, read_only=True)
    try:
        linhas = list(workbook["Arquivos"].iter_rows(min_row=2, values_only=True))
        resumo_valores = [
            valor
            for linha in workbook["Resumo"].iter_rows(values_only=True)
            for valor in linha
            if valor
        ]
    finally:
        workbook.close()

    assert len(linhas) == 5  # O inventário completo é preservado.
    assert any("CANCELADA PELO USUÁRIO" in str(valor) for valor in resumo_valores)
    assert any("NÃO PROCESSADO" in str(linha[9]) for linha in linhas)


def test_cancelamento_antes_de_iniciar(sample_tree: Path, tmp_path: Path) -> None:
    """Cancelar antes do início encerra a análise sem processar arquivos."""
    engine = AnalysisEngine(_config(sample_tree, tmp_path))
    engine.cancel()
    resumo = engine.run()

    assert resumo.cancelled
    assert resumo.files_processed == 0
    assert resumo.report_path is not None and resumo.report_path.exists()


def test_analise_em_thread_separada(sample_tree: Path, tmp_path: Path) -> None:
    """O motor pode ser executado fora da thread principal sem travar."""
    engine = AnalysisEngine(_config(sample_tree, tmp_path))
    resultado: list[object] = []

    thread = threading.Thread(target=lambda: resultado.append(engine.run()))
    thread.start()
    thread.join(timeout=120)

    assert not thread.is_alive()
    assert resultado and getattr(resultado[0], "files_processed", 0) == 5


def test_pasta_vazia(tmp_path: Path) -> None:
    """Uma origem sem arquivos gera relatório válido e vazio."""
    origem = tmp_path / "vazia"
    origem.mkdir()
    resumo = AnalysisEngine(_config(origem, tmp_path)).run()

    assert resumo.files_found == 0
    assert resumo.report_path is not None and resumo.report_path.exists()


def test_nao_altera_os_arquivos_analisados(sample_tree: Path, tmp_path: Path) -> None:
    """Os arquivos analisados não podem ter conteúdo ou datas alterados."""
    antes = {
        caminho: (caminho.stat().st_size, caminho.stat().st_mtime_ns, caminho.read_bytes())
        for caminho in sample_tree.rglob("*")
        if caminho.is_file()
    }

    AnalysisEngine(_config(sample_tree, tmp_path)).run()

    depois = {
        caminho: (caminho.stat().st_size, caminho.stat().st_mtime_ns, caminho.read_bytes())
        for caminho in sample_tree.rglob("*")
        if caminho.is_file()
    }
    assert antes == depois


def test_arquivo_grande_processado_em_blocos(tmp_path: Path) -> None:
    """Arquivo de vários blocos é processado corretamente pelo motor."""
    import hashlib

    origem = tmp_path / "grande"
    origem.mkdir()
    conteudo = os.urandom(1024) * 6000  # ~6 MB
    (origem / "grande.bin").write_bytes(conteudo)

    config = _config(origem, tmp_path, chunk_size=1024 * 1024)
    resumo = AnalysisEngine(config).run()

    workbook = load_workbook(resumo.report_path, read_only=True)
    try:
        linha = next(workbook["Arquivos"].iter_rows(min_row=2, max_row=2, values_only=True))
    finally:
        workbook.close()

    assert linha[8] == hashlib.sha256(conteudo).hexdigest().upper()
    assert linha[5] == len(conteudo)
